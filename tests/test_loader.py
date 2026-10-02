from pathlib import Path

import pytest

from app.ingestion.loaders import (
    DocumentLoaderError,
    DocumentPage,
    load_document,
)


def test_load_txt_document(tmp_path: Path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text(
        "Atlas RAG is an enterprise knowledge assistant.",
        encoding="utf-8",
    )

    pages = load_document(str(file_path))

    assert len(pages) == 1
    assert isinstance(pages[0], DocumentPage)
    assert pages[0].text == (
        "Atlas RAG is an enterprise knowledge assistant."
    )
    assert pages[0].page_number is None


def test_load_markdown_document(tmp_path: Path):
    file_path = tmp_path / "sample.md"
    file_path.write_text(
        "# Atlas RAG\n\nEnterprise knowledge assistant.",
        encoding="utf-8",
    )

    pages = load_document(str(file_path))

    assert len(pages) == 1
    assert pages[0].text == (
        "# Atlas RAG\n\nEnterprise knowledge assistant."
    )
    assert pages[0].page_number is None


def test_load_docx_document(tmp_path: Path):
    from docx import Document

    file_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("Atlas RAG")
    document.add_paragraph("Enterprise knowledge assistant.")
    document.save(file_path)

    pages = load_document(str(file_path))

    assert len(pages) == 1
    assert pages[0].text == (
        "Atlas RAG\n\nEnterprise knowledge assistant."
    )
    assert pages[0].page_number is None


def test_missing_file_raises_error():
    with pytest.raises(DocumentLoaderError, match="File not found"):
        load_document("tests/does-not-exist.txt")


def test_unsupported_file_type_raises_error(tmp_path: Path):
    file_path = tmp_path / "sample.exe"
    file_path.write_bytes(b"fake binary content")

    with pytest.raises(
        DocumentLoaderError,
        match="Unsupported file type",
    ):
        load_document(str(file_path))


def test_empty_text_file_returns_empty_list(tmp_path: Path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("", encoding="utf-8")

    pages = load_document(str(file_path))

    assert pages == []


def test_whitespace_only_text_file_returns_empty_list(tmp_path: Path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("   \n\n   ", encoding="utf-8")

    pages = load_document(str(file_path))

    assert pages == []
